import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from app.utils import iso_now


def build_export_xlsx(tasks):
    wb = Workbook()
    ws = wb.active
    ws.title = "Tasks"
    FONT = "Arial"
    headers = ["ID", "Type", "Title", "Status", "Priority", "Project(s)", "Tags",
               "Discussed With", "Discussed", "Start", "Due", "Done", "Closed",
               "Notes", "Cancel Reason", "Updated"]

    header_fill = PatternFill("solid", fgColor="4D6D8C")
    header_font = Font(name=FONT, size=10, bold=True, color="FFFFFF")
    thin = Side(style="thin", color="E3DDD0")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    status_fill = {
        "Done": PatternFill("solid", fgColor="E6F0E9"),
        "In Progress": PatternFill("solid", fgColor="E7EDF3"),
        "In Review": PatternFill("solid", fgColor="EFE7F5"),
        "Pending": PatternFill("solid", fgColor="F6ECD9"),
        "To Do": PatternFill("solid", fgColor="EFE9DC"),
        "Cancelled": PatternFill("solid", fgColor="F5E2E0"),
    }

    for c, h in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=c, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.border = border
        cell.alignment = Alignment(vertical="center")
    ws.freeze_panes = "A2"

    rows_sorted = sorted(tasks, key=lambda t: (t.get("updated_ts") or t.get("updated_at") or ""), reverse=True)
    row = 2
    for t in rows_sorted:
        discussed = t.get("discussed_from") or ""
        if t.get("discussed_to") and t.get("discussed_to") != discussed:
            discussed = f"{discussed} - {t['discussed_to']}"
        values = [
            t.get("id", ""), t.get("type") or "Task", t.get("title", ""), t.get("status", ""),
            t.get("priority") or "P3", ", ".join(t.get("project") or []), ", ".join(t.get("tags") or []),
            ", ".join(t.get("discussed_with") or []), discussed, t.get("start_date") or "",
            t.get("due_date") or "", t.get("done_at") or "", t.get("closed_at") or "",
            t.get("notes", ""), t.get("cancel_reason", ""), t.get("updated_at", ""),
        ]
        for c, v in enumerate(values, start=1):
            cell = ws.cell(row=row, column=c, value=v)
            cell.font = Font(name=FONT, size=10)
            cell.alignment = Alignment(vertical="top", wrap_text=(c in (3, 14)))
            cell.border = border
        status_cell = ws.cell(row=row, column=4)
        status_cell.fill = status_fill.get(t.get("status", ""), PatternFill())
        row += 1

    last_row = max(row - 1, 1)
    widths = {1: 7, 2: 9, 3: 52, 4: 12, 5: 9, 6: 24, 7: 16, 8: 20, 9: 18, 10: 11, 11: 11, 12: 11, 13: 11, 14: 46, 15: 24, 16: 11}
    for c, w in widths.items():
        ws.column_dimensions[get_column_letter(c)].width = w

    if last_row >= 1:
        table = Table(displayName="ExportedTasks", ref=f"A1:P{last_row}")
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
        ws.add_table(table)

    ws2 = wb.create_sheet("Summary")
    ws2["A1"] = "Exported"
    ws2["B1"] = iso_now()
    ws2["A2"] = "Task count"
    ws2["B2"] = len(tasks)
    for cell in (ws2["A1"], ws2["B1"], ws2["A2"], ws2["B2"]):
        cell.font = Font(name=FONT, size=10)
    ws2.column_dimensions["A"].width = 14
    ws2.column_dimensions["B"].width = 24

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
